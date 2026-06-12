from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_PATH = Path("docs/commissioning_agv_factory_acceptance_checklist.docx")

ACCENT = RGBColor(31, 78, 121)
LIGHT_FILL = "DCE6F1"
HEADER_FILL = "BDD7EE"
BORDER_COLOR = "A6A6A6"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = tc_mar.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER_COLOR, size=8):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def format_run(run, *, bold=False, size=11, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if color is not None:
        run.font.color.rgb = color


def add_paragraph(doc, text="", *, style=None, align=None, bold=False, size=11, color=None, space_after=6):
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    format_run(run, bold=bold, size=size, color=color)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.15
    return paragraph


def add_cell_text(cell, text, *, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    format_run(run, bold=bold, size=size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def configure_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("CHECKLIST COMMISSIONING AGV")
    format_run(run, bold=True, size=9, color=ACCENT)
    p.border_bottom = True

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    run = fp.add_run("Halaman ")
    format_run(run, size=9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = fp.add_run()
    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_end)
    format_run(r, size=9)


def add_metadata_table(doc):
    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.4), Cm(6.2), Cm(3.4), Cm(6.2)]
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = width

    rows = [
        ("Jenis Dokumen", "Commissioning / Factory Acceptance Checklist", "Nomor Dokumen", "................................"),
        ("Nama Unit", "AGV KIM2A", "Nomor Seri", "................................"),
        ("Lokasi Uji", "Factory / Workshop", "Tanggal Uji", "................................"),
        ("Customer", "Generic Customer", "Diperiksa Oleh", "................................"),
    ]
    for row, values in zip(table.rows, rows):
        for idx, value in enumerate(values):
            is_label = idx % 2 == 0
            add_cell_text(row.cells[idx], value, bold=is_label)
            if is_label:
                set_cell_shading(row.cells[idx], LIGHT_FILL)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_info_box(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0].cells[0].width = Cm(17.0)
    cell = table.cell(0, 0)
    set_table_borders(table, color="9EADBA")
    set_cell_shading(cell, "F7FBFF")
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(3)
    r1 = p1.add_run(title)
    format_run(r1, bold=True, size=11, color=ACCENT)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(text)
    format_run(r2, size=10.5)
    set_cell_margins(cell, top=110, start=150, bottom=110, end=150)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_checklist_table(doc, title, items):
    add_paragraph(doc, title, style="Heading 2", bold=True, size=13, color=ACCENT, space_after=6)
    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["No", "Item Pemeriksaan", "Kriteria / Metode Verifikasi", "OK", "NG", "N/A", "Catatan"]
    widths = [Cm(1.0), Cm(5.1), Cm(7.2), Cm(1.0), Cm(1.0), Cm(1.1), Cm(3.6)]
    for i, width in enumerate(widths):
        table.rows[0].cells[i].width = width
    for cell, header in zip(table.rows[0].cells, headers):
        add_cell_text(cell, header, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, HEADER_FILL)
    set_repeat_table_header(table.rows[0])

    for idx, item in enumerate(items, start=1):
        row = table.add_row()
        values = [str(idx), item[0], item[1], "", "", "", ""]
        aligns = [
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
        ]
        for cell, value, align in zip(row.cells, values, aligns):
            add_cell_text(cell, value, size=10, align=align)

    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def build_document():
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    for name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 13)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    section = doc.sections[0]
    configure_page(section)
    add_header_footer(section)

    add_paragraph(
        doc,
        "CHECKLIST COMMISSIONING AGV",
        style="Title",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=22,
        color=ACCENT,
        space_after=4,
    )
    add_paragraph(
        doc,
        "Pemeriksaan unit oleh customer di factory sebelum delivery",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=11,
        space_after=14,
    )

    add_metadata_table(doc)

    add_paragraph(doc, "1. Tujuan Dokumen", style="Heading 1", bold=True, size=16, color=ACCENT, space_after=6)
    add_paragraph(
        doc,
        "Dokumen ini digunakan sebagai checklist commissioning untuk verifikasi kondisi, fungsi, dan keselamatan unit AGV sebelum diserahkan ke customer. Seluruh item diperiksa di factory dengan metode inspeksi visual, functional test, dan witness test sesuai kebutuhan.",
        size=11,
        space_after=8,
    )

    add_paragraph(doc, "2. Ruang Lingkup Unit", style="Heading 1", bold=True, size=16, color=ACCENT, space_after=6)
    add_info_box(
        doc,
        "Deskripsi Singkat Sistem",
        "AGV merupakan differential-drive automated vehicle dengan mode manual dan auto tape following. Sistem menggunakan sensor magnetic guide, pembacaan RFID, discrete I/O, analog speed output, serta safety watchdog untuk memastikan unit berhenti aman saat terjadi fault, tape loss, atau emergency stop. Dokumen ini tidak mencakup feature SLMP, PLC handshake, maupun fungsi apa pun yang terkait trolley.",
    )

    add_paragraph(doc, "3. Operating Description", style="Heading 1", bold=True, size=16, color=ACCENT, space_after=6)
    operating_points = [
        "Manual mode: operator menjalankan unit melalui pendant/jog input untuk gerakan forward, reverse, left, right, dan kombinasi steering.",
        "Auto mode: unit mengikuti magnetic tape secara otomatis dengan kontrol kecepatan dan koreksi arah berbasis sensor guide.",
        "RFID dan marker sequence: unit dapat merespons tag RFID dan marker sebagai trigger logika operasi yang telah diprogram pada profile AGV.",
        "Safety response: saat emergency aktif, sensor timeout terjadi, atau tape hilang, unit harus melakukan braking dan menghentikan gerakan secara aman.",
        "Reverse auto: unit dapat melakukan reverse tape following bila fungsi ini di-enable pada pengujian.",
        "Pusher otomatis: mekanisme pusher bergerak naik dan turun secara otomatis sesuai urutan operasi yang ditentukan, dengan interlock posisi dan respon gerak yang stabil.",
    ]
    for point in operating_points:
        add_paragraph(doc, point, size=10.8, space_after=4)

    visual_items = [
        ("Kondisi fisik unit", "Body, cover, rangka, bumper, dan panel tidak rusak, tidak longgar, dan finishing rapi."),
        ("Nameplate dan label", "Label identifikasi unit, arah gerak, warning, dan emergency stop terpasang jelas dan terbaca."),
        ("Kerapian wiring", "Wiring internal dan eksternal rapi, terlindungi, dan tidak ada kabel terjepit atau terkelupas."),
        ("Kondisi roda dan mekanik", "Roda penggerak, caster, brake linkage, dan fastener dalam kondisi baik."),
        ("Kondisi pusher", "Mekanisme pusher, guide, actuator, dan stopper terpasang kuat serta bergerak bebas tanpa gesekan abnormal."),
    ]
    add_checklist_table(doc, "4. Checklist Inspeksi Visual dan Mekanik", visual_items)

    safety_items = [
        ("Emergency stop", "Saat emergency stop ditekan, gerak AGV berhenti aman dan status fault muncul sesuai desain."),
        ("Recovery emergency", "Setelah reset sesuai prosedur, unit dapat kembali ke kondisi siap tanpa alarm tersisa."),
        ("Brake saat tape loss", "Pada kondisi tape hilang, unit melakukan braking dan tidak melanjutkan gerak liar."),
        ("Brake saat communication fault", "Saat watchdog mendeteksi fault komunikasi sensor utama, output gerak turun ke kondisi aman."),
        ("Mode transition safety", "Perpindahan manual/auto tidak menimbulkan gerakan mendadak atau tidak terkendali."),
        ("Pusher safety interlock", "Pusher tidak bergerak di luar urutan operasi yang diizinkan dan berhenti aman saat fault/emergency."),
    ]
    add_checklist_table(doc, "5. Checklist Safety Function", safety_items)

    manual_items = [
        ("Manual forward", "Unit bergerak maju stabil saat perintah forward diberikan."),
        ("Manual reverse", "Unit bergerak mundur stabil saat perintah reverse diberikan."),
        ("Manual left / right", "Unit merespons perintah belok kiri dan kanan dengan arah yang benar."),
        ("Manual combined steering", "Kombinasi forward-left, forward-right, reverse-left, dan reverse-right bekerja sesuai logika."),
        ("Brake / idle command", "Saat tidak ada perintah jog, AGV kembali ke kondisi brake atau idle sesuai desain."),
    ]
    add_checklist_table(doc, "6. Checklist Fungsi Manual Mode", manual_items)

    auto_items = [
        ("Start auto", "Unit dapat masuk ke auto run dengan pre-condition yang benar dan tape terdeteksi."),
        ("Tape following stability", "AGV mengikuti jalur tape secara stabil tanpa osilasi berlebihan."),
        ("Speed mode response", "Perubahan speed mode HIGH / SLOW / EXTRA_SLOW bekerja sesuai sequence dan kondisi operasi."),
        ("Marker detection", "Marker kiri/kanan terbaca dan memicu aksi yang sesuai."),
        ("RFID detection", "Tag RFID terbaca konsisten dan dapat digunakan sebagai trigger sequence."),
        ("Sequence stop / resume", "Saat sequence meminta stop, AGV berhenti sesuai logika dan dapat resume dengan benar."),
        ("Reverse auto", "Bila diuji, fungsi reverse auto mengikuti tape secara terkendali pada kecepatan yang diizinkan."),
    ]
    add_checklist_table(doc, "7. Checklist Fungsi Auto dan Navigasi", auto_items)

    pusher_items = [
        ("Deskripsi operasi pusher tersedia", "Urutan kerja pusher naik dan turun dijelaskan pada dokumen dan dipahami pihak penguji."),
        ("Auto up command", "Pusher bergerak naik otomatis saat trigger operasi diberikan."),
        ("Auto down command", "Pusher bergerak turun otomatis saat trigger operasi diberikan."),
        ("Posisi akhir up", "Posisi akhir saat pusher up tercapai konsisten dan tidak overshoot."),
        ("Posisi akhir down", "Posisi akhir saat pusher down tercapai konsisten dan tidak macet."),
        ("Cycle repeatability", "Minimal tiga siklus up/down berturut-turut berjalan normal tanpa fault."),
        ("Interlock dengan AGV operation", "Pergerakan pusher tidak mengganggu stabilitas AGV dan mengikuti urutan proses yang diizinkan."),
    ]
    add_checklist_table(doc, "8. Checklist Fitur Pusher Otomatis", pusher_items)

    performance_items = [
        ("Arah putaran motor", "Arah motor kiri dan kanan sesuai command dan tidak tertukar."),
        ("Respons acceleration / deceleration", "Percepatan dan perlambatan terasa halus serta tidak menimbulkan hentakan abnormal."),
        ("Stabilitas output analog speed", "Output kecepatan analog berubah konsisten terhadap command operasi."),
        ("Keseimbangan gerak", "Tidak ada tarikan berlebih ke satu sisi saat lintasan lurus dan kondisi tape normal."),
        ("Kebisingan / getaran", "Tidak ada bunyi atau getaran abnormal selama pengujian commissioning."),
    ]
    add_checklist_table(doc, "9. Checklist Performa Operasi", performance_items)

    add_paragraph(doc, "10. Catatan Commissioning", style="Heading 1", bold=True, size=16, color=ACCENT, space_after=6)
    note_table = doc.add_table(rows=4, cols=1)
    note_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    note_table.autofit = False
    for row in note_table.rows:
        row.cells[0].width = Cm(17.3)
        add_cell_text(row.cells[0], "........................................................................................................................", size=10.5)
    set_table_borders(note_table)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_paragraph(doc, "11. Hasil Akhir dan Sign-Off", style="Heading 1", bold=True, size=16, color=ACCENT, space_after=6)
    sign_table = doc.add_table(rows=4, cols=4)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.autofit = False
    sign_widths = [Cm(4.0), Cm(4.6), Cm(4.0), Cm(4.7)]
    for row in sign_table.rows:
        for i, width in enumerate(sign_widths):
            row.cells[i].width = width
    sign_rows = [
        ("Status Commissioning", "PASS / HOLD / FAIL", "Tanggal", "................................"),
        ("Disiapkan Oleh", "................................", "Diverifikasi Oleh", "................................"),
        ("Jabatan", "................................", "Jabatan", "................................"),
        ("Tanda Tangan", "\n\n................................", "Tanda Tangan", "\n\n................................"),
    ]
    for row, values in zip(sign_table.rows, sign_rows):
        for idx, value in enumerate(values):
            add_cell_text(row.cells[idx], value, bold=idx % 2 == 0)
            if idx % 2 == 0:
                set_cell_shading(row.cells[idx], LIGHT_FILL)
    set_table_borders(sign_table)

    return doc


def main():
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT_PATH)
    print(f"Created {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
